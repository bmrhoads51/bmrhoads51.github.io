"""Build the GNN article page from the newest supplied Word manuscript."""
from __future__ import annotations

from html import escape
from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / "GNN" / "GNNPaper_.docx"
ASSETS = ROOT / "assets" / "gnn"
ASSETS.mkdir(parents=True, exist_ok=True)
DOC = Document(SOURCE)

R_EMBED = qn("r:embed")
W_VAL = qn("w:val")


def math_text(node):
    """Render common Word equation structures as readable inline HTML."""
    name = node.tag.split("}")[-1]
    children = list(node)
    if name == "t":
        return escape(node.text or "")
    if name == "f":
        num = next((x for x in children if x.tag.endswith("}num")), None)
        den = next((x for x in children if x.tag.endswith("}den")), None)
        return f'<span class="fraction"><span>{math_text(num) if num is not None else ""}</span><span>{math_text(den) if den is not None else ""}</span></span>'
    if name in ("sSup", "sSub", "sSubSup"):
        base = next((x for x in children if x.tag.endswith("}e")), None)
        sub = next((x for x in children if x.tag.endswith("}sub")), None)
        sup = next((x for x in children if x.tag.endswith("}sup")), None)
        return (math_text(base) if base is not None else "") + (f"<sub>{math_text(sub)}</sub>" if sub is not None else "") + (f"<sup>{math_text(sup)}</sup>" if sup is not None else "")
    if name == "rad":
        deg = next((x for x in children if x.tag.endswith("}deg")), None)
        value = next((x for x in children if x.tag.endswith("}e")), None)
        return f'<span class="radical">√<span>{math_text(value) if value is not None else ""}</span></span>' if deg is None or not math_text(deg) else f'<span class="radical"><sup>{math_text(deg)}</sup>√<span>{math_text(value) if value is not None else ""}</span></span>'
    if name == "d":
        props = next((x for x in children if x.tag.endswith("}dPr")), None)
        begin, end = "(", ")"
        if props is not None:
            b = next((x for x in props if x.tag.endswith("}begChr")), None)
            e = next((x for x in props if x.tag.endswith("}endChr")), None)
            if b is not None: begin = b.get(W_VAL, begin)
            if e is not None: end = e.get(W_VAL, end)
        return escape(begin) + "".join(math_text(x) for x in children if not x.tag.endswith("}dPr")) + escape(end)
    if name in ("oMath", "oMathPara", "e", "num", "den", "sub", "sup", "r", "func", "fName", "nary", "limLow", "limUpp", "acc", "box", "bar", "groupChr"):
        return "".join(math_text(x) for x in children if not x.tag.endswith("Pr"))
    if name == "brk":
        return " "
    if name.endswith("Pr") or name in ("rPr", "ctrlPr"):
        return ""
    return "".join(math_text(x) for x in children)


def run_html(run):
    pieces = []
    for child in run._r:
        tag = child.tag.split("}")[-1]
        if tag == "t": pieces.append(escape(child.text or ""))
        elif tag in ("tab",): pieces.append(" ")
        elif tag in ("br", "cr"): pieces.append("<br>")
        elif tag == "oMath": pieces.append(f'<span class="math">{math_text(child)}</span>')
    value = "".join(pieces)
    if run.bold: value = f"<strong>{value}</strong>"
    if run.italic: value = f"<em>{value}</em>"
    if run.underline: value = f"<u>{value}</u>"
    return value


def paragraph_html(p):
    output = []
    for child in p._p:
        tag = child.tag.split("}")[-1]
        if tag == "r": output.append(run_html(Run(child, p)))
        elif tag in ("oMath", "oMathPara"): output.append(f'<span class="math">{math_text(child)}</span>')
        elif tag == "hyperlink":
            for run in child:
                if run.tag.endswith("}r"):
                    output.append(run_html(Run(run, p)))
    rendered = "".join(output).strip()
    return re.sub(r'#\s*(Equation\s+\d+)', r'<span class="equation-number">\1</span>', rendered)


def images_in(p):
    images = []
    for blip in p._p.xpath('.//a:blip'):
        rid = blip.get(R_EMBED)
        if not rid: continue
        part = p.part.related_parts[rid]
        filename = Path(part.partname).name
        target = ASSETS / filename
        target.write_bytes(part.blob)
        images.append(f"assets/gnn/{filename}")
    return images


def table_html(table, caption):
    rows = []
    for ri, row in enumerate(table.rows):
        cells = []
        for cell in row.cells:
            content = "<br>".join(escape(p.text) for p in cell.paragraphs if p.text.strip())
            element = "th" if ri == 0 else "td"
            cells.append(f"<{element}>{content}</{element}>")
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return f'<div class="table-scroll"><table><caption>{caption}</caption><tbody>{"".join(rows)}</tbody></table></div>'


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def figure_tag(image):
    figure_ids = {
        "image1.png": "figure-2-1",
        "image2.png": "figure-2-2",
        "image3.png": "figure-2-3",
    }
    ident = figure_ids.get(Path(image).name)
    return f'<figure id="{ident}">' if ident else "<figure>"


MAJOR = {"Introduction", "Methods", "Results", "Discussion", "Conclusion", "Supplemental", "References"}
CONTENT = []
TOC = []
pending_figure = None
pending_table_caption = ""
in_references = False
ref_items = []
heading_id = set()


def add_heading(label, level):
    ident = slug(label)
    if ident in heading_id: ident += f"-{len(heading_id)}"
    heading_id.add(ident)
    CONTENT.append(f'<h{level} id="{ident}">{escape(label)}</h{level}>')
    if level == 2: TOC.append((label, ident))


for child in DOC.element.body.iterchildren():
    tag = child.tag.split("}")[-1]
    if tag == "tbl":
        CONTENT.append(table_html(Table(child, DOC), pending_table_caption))
        pending_table_caption = ""
        continue
    if tag != "p": continue
    p = Paragraph(child, DOC)
    raw = p.text.strip()
    html = paragraph_html(p)
    imgs = images_in(p)
    if raw in ("Property Prediction of NiAl Microstructure using Graph Neural Networks", "Structure-Property Linkage in Alloys using Graph Neural Network and Explainable Artificial Intelligence"):
        continue
    if p.style.name == "EndNote Bibliography":
        in_references = True
        ref_items.append(re.sub(r"^\d+\.\s*", "", html))
        continue
    if in_references and raw and raw != "References":
        in_references = False
    if not raw and not imgs and not p._p.xpath('.//m:oMath'): continue
    if raw.startswith("Table 1") or raw.startswith("Table 2"):
        if CONTENT and CONTENT[-1].startswith('<div class="table-scroll"><table><caption></caption>'):
            CONTENT[-1] = CONTENT[-1].replace('<caption></caption>', f'<caption>{html}</caption>', 1)
        else:
            pending_table_caption = html
        continue
    if raw.lower().startswith("abstract:"):
        abstract_body = re.sub(r'^(?:<strong>)?Abstract:\s*(?:</strong>)?', '', html, flags=re.I)
        CONTENT.append(f'<section class="abstract"><h2 id="abstract">Abstract</h2><p>{abstract_body.strip()}</p></section>')
        TOC.append(("Abstract", "abstract"))
        continue
    clean = raw.rstrip(": ")
    if clean in MAJOR:
        if pending_figure: CONTENT.append(pending_figure); pending_figure = None
        add_heading(clean, 2)
        if clean == "References": in_references = True
        continue
    if re.match(r"^(?:2|3)\.\d+\.?\s", clean) or re.match(r"^S1\.\d+", clean):
        add_heading(clean, 3)
        continue
    if p.style.name == "List Paragraph" and raw and not raw.startswith("Figure"):
        add_heading(clean, 4)
        if imgs:
            pending_figure = f'{figure_tag(imgs[0])}<img src="{imgs[0]}" alt="Graph neural network diagram from the paper" loading="lazy">'
        continue
    if imgs:
        if pending_figure and raw.startswith("Figure"):
            CONTENT.append(pending_figure + f'<figcaption>{html}</figcaption></figure>')
            pending_figure = f'{figure_tag(imgs[0])}<img src="{imgs[0]}" alt="Research figure from the manuscript" loading="lazy">'
            continue
        if pending_figure: CONTENT.append(pending_figure + "</figure>")
        pending_figure = f'{figure_tag(imgs[0])}<img src="{imgs[0]}" alt="{escape(raw[:130] or "Research figure from the manuscript", quote=True)}" loading="lazy">'
        if raw.startswith("Figure"):
            CONTENT.append(pending_figure + f'<figcaption>{html}</figcaption></figure>')
            pending_figure = None
        continue
    if raw.startswith("Figure") and pending_figure:
        CONTENT.append(pending_figure + f'<figcaption>{html}</figcaption></figure>')
        pending_figure = None
        continue
    if pending_figure:
        CONTENT.append(pending_figure + "</figure>")
        pending_figure = None
    if raw.startswith("Figure"):
        CONTENT.append(f'<p class="caption">{html}</p>')
    elif not raw and html:
        CONTENT.append(f'<div class="equation">{html}</div>')
    elif raw:
        CONTENT.append(f"<p>{html}</p>")

if pending_figure: CONTENT.append(pending_figure + "</figure>")
if ref_items:
    CONTENT.append('<ol class="references">' + "".join(f"<li>{x}</li>" for x in ref_items) + "</ol>")

toc = "".join(f'<a href="#{ident}">{escape(label)}</a>' for label, ident in TOC)
article = f'''<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
  <meta name="description" content="Full research manuscript on graph neural networks, Ni–Al microstructures, and explainable artificial intelligence by Benjamin Rhoads and collaborators.">
  <title>Structure–Property Linkage in Alloys | Benjamin Rhoads</title>
  <link rel="stylesheet" href="style.css">
</head>
<body>
  <a class="skip-link" href="#main">Skip to content</a>
  <header class="site-header"><div class="wrap header-inner"><a class="brand" href="index.html">Benjamin Rhoads</a><nav aria-label="Main navigation"><a href="index.html">Home</a><a href="publications.html">Publications</a><a href="projects.html">Projects</a></nav></div></header>
  <main id="main" class="wrap paper-layout">
    <aside class="paper-sidebar" aria-label="Article navigation"><p class="eyebrow">On this page</p>{toc}</aside>
    <article class="paper">
      <div class="paper-top"><p class="eyebrow">Publication · Materials · 2025</p><h1>Structure–Property Linkage in Alloys Using Graph Neural Network and Explainable Artificial Intelligence</h1><p class="authors">Benjamin Rhoads · Abigail Hogue · Lars Kotthoff · Samrat Choudhury</p><p class="journal"><em>Materials</em> 18(16), 3778 · DOI: 10.3390/ma18163778</p><div class="actions"><a class="button" href="https://doi.org/10.3390/ma18163778" target="_blank" rel="noopener noreferrer">Read published article ↗</a></div></div>
      {chr(10).join(CONTENT)}
    </article>
  </main>
  <footer class="site-footer"><div class="wrap">© Benjamin Rhoads · <a href="index.html">Home</a> · <a href="https://github.com/bmrhoads51" target="_blank" rel="noopener noreferrer">GitHub</a></div></footer>
</body>
</html>'''
(ROOT / "gnn-paper.html").write_text(article, encoding="utf-8")
print(f"Wrote gnn-paper.html, {len(CONTENT)} content blocks, {len(ref_items)} references, {len(list(ASSETS.glob('*')))} images")
